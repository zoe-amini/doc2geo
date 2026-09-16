"""Word documents.

Coordinates in a .docx live in one of two places: a table of survey points, or a sentence
in the middle of a paragraph. Both are pulled out here and handed on separately.
"""

from __future__ import annotations

from pathlib import Path

from ..records import Extraction, Table, TextBlock


def read_docx(path: Path) -> Extraction:
    import docx  # python-docx

    document = docx.Document(str(path))
    tables: list[Table] = []
    for index, table in enumerate(document.tables, 1):
        rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
        rows = [r for r in rows if any(r)]
        if not rows:
            continue
        tables.append(Table(header=rows[0], rows=rows[1:], name=f"table_{index}"))

    paragraphs = [p.text.strip() for p in document.paragraphs if p.text.strip()]
    blocks = [TextBlock(text="\n".join(paragraphs), backend="python-docx")] if paragraphs else []
    return Extraction(source=path, tables=tables, blocks=blocks, pages=1, backend="python-docx")
