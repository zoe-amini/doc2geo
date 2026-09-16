"""CSV, TSV and Excel workbooks.

The easy case, and the one that carries the most rows in practice. Each sheet becomes one
`Table`; the first non-empty row is taken as the header when it looks like labels rather
than data.
"""

from __future__ import annotations

import csv
from pathlib import Path

from ..records import Extraction, Table

CSV_SUFFIXES = {".csv", ".tsv", ".txt"}
EXCEL_SUFFIXES = {".xlsx", ".xlsm"}


def _looks_like_header(row: list[str]) -> bool:
    """A header row is mostly non-numeric text."""
    cells = [c for c in row if str(c).strip()]
    if not cells:
        return False
    numeric = 0
    for cell in cells:
        try:
            float(str(cell).replace(",", "").replace("°", ""))
            numeric += 1
        except ValueError:
            pass
    return numeric <= len(cells) / 2


def _to_table(rows: list[list[str]], name: str) -> Table | None:
    rows = [r for r in rows if any(str(c).strip() for c in r)]
    if not rows:
        return None
    if _looks_like_header(rows[0]):
        return Table(header=[str(c).strip() for c in rows[0]], rows=rows[1:], name=name)
    width = max(len(r) for r in rows)
    return Table(header=[f"col_{i + 1}" for i in range(width)], rows=rows, name=name)


def read_csv(path: Path) -> Extraction:
    text = path.read_text(encoding="utf-8-sig", errors="replace")
    sample = text[:8192]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel_tab if path.suffix.lower() == ".tsv" else csv.excel
    rows = [[str(c) for c in row] for row in csv.reader(text.splitlines(), dialect)]
    table = _to_table(rows, path.stem)
    return Extraction(source=path, tables=[table] if table else [], pages=1, backend="csv")


def read_excel(path: Path) -> Extraction:
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    tables: list[Table] = []
    try:
        for index, sheet in enumerate(workbook.worksheets, 1):
            rows = [
                ["" if cell is None else str(cell) for cell in row]
                for row in sheet.iter_rows(values_only=True)
            ]
            table = _to_table(rows, sheet.title)
            if table:
                table.page = index
                tables.append(table)
    finally:
        workbook.close()
    return Extraction(source=path, tables=tables, pages=len(tables), backend="openpyxl")
